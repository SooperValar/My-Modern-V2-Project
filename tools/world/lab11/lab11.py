#!/usr/bin/env python3
"""Lab 11, the Sunken Dark - the layout lives in an Excel sheet (lab11_layout.xlsx).

    lab11.py export <xlsx>             write the sheet from the original layout (draft_layout.py)
    lab11.py build  <xlsx> <out>       read the sheet: checks, preview.png, lab11.spec
Apply with:  build_lab <.dat folder> lab11.spec ; then "server pop" and "server light".

Sheet "Map": one cell per tile, column header = map x, row header = map y.
Sheet "Sections": the four flooded sections (each one's water rectangle drains together).
Codes (not case-sensitive):
  #  wall                ~  water (drains)       .  dry floor (stone)    e  hall floor (no monsters)
  T  tavern floor (log out)   ct counter          |  bookshelf (blocks)   s  bookshelf, rotten pages
  S1-S4 bookshelf with section 1-4's tome         L1-L4 lectern           D1-D4 gate
  B  brazier (players light it)   C  lit bowl     m  monster (kind from the section it is in)
  E1-E3 exploration points        Q  shrine of Kwai   H  offering shrine
  X  entry/exit portal   A  arrival (where the lab gate drops you)        V  end portal (vault)
  BAR barkeeper   MW weapon merchant   MA armor merchant
  SW/SF/SC/SD  the standard shrine-of-Kwai room: red brick walls, floor, black candles, Seyan'Du door
Objects that stand in the water (|, s, S, E, Q) sit on water when two or more of their four
neighbours are water, otherwise on dry floor."""
import json, random, sys, os
from collections import deque

OX, OY = 180, 900            # map position of the sheet's top-left tile
W, H = 120, 100
TICKS = 18
CODES = {   # code: (meaning, fill colour, text colour)
    '#': ('wall', '2B2B30', 'FFFFFF'),
    '~': ('water (drains when its lectern gets the tome)', '2E6496', 'FFFFFF'),
    '.': ('dry floor', 'A39E92', '000000'),
    'e': ('hall floor, no monsters', '8C5A4B', 'FFFFFF'),
    'T': ('tavern floor - stepping on it logs you out', 'D08A4E', '000000'),
    'ct': ('bar counter', 'E3C531', '000000'),
    '|': ('bookshelf (blocks)', '7A5232', 'FFFFFF'),
    's': ('bookshelf with only rotten pages', '3D7A47', 'FFFFFF'),
    'S1': ('bookshelf with tome I (section 1)', '3CE65A', '000000'),
    'S2': ('bookshelf with tome II (section 2)', '3CE65A', '000000'),
    'S3': ('bookshelf with tome III (section 3)', '3CE65A', '000000'),
    'S4': ('bookshelf with tome IV (section 4)', '3CE65A', '000000'),
    'L1': ('lectern of section 1', 'E63CE6', '000000'),
    'L2': ('lectern of section 2', 'E63CE6', '000000'),
    'L3': ('lectern of section 3', 'E63CE6', '000000'),
    'L4': ('lectern of section 4', 'E63CE6', '000000'),
    'D1': ('gate opened by lectern 1', 'C87828', '000000'),
    'D2': ('gate opened by lectern 2', 'C87828', '000000'),
    'D3': ('gate opened by lectern 3', 'C87828', '000000'),
    'D4': ('gate opened by lectern 4', 'C87828', '000000'),
    'B': ('brazier - cold until a player lights it (5 min)', 'FFA028', '000000'),
    'C': ('lit bowl / candle (always burning)', 'FFE078', '000000'),
    'm': ('monster', 'C81E1E', 'FFFFFF'),
    'E1': ('exploration point 1', '00FFFF', '000000'),
    'E2': ('exploration point 2', '00FFFF', '000000'),
    'E3': ('exploration point 3', '00FFFF', '000000'),
    'Q': ('shrine of Kwai', 'FFFFFF', '000000'),
    'SD': ("shrine room: Seyan'Du door (only Seyan'Du open it)", '9B30FF', 'FFFFFF'),
    'SW': ('shrine room: red brick wall', '8B2A1E', 'FFFFFF'),
    'SF': ('shrine room: floor', 'B07060', '000000'),
    'SC': ('shrine room: black candle', '503050', 'FFFFFF'),
    'H': ('offering shrine', 'F0DC3C', '000000'),
    'X': ('entry/exit portal', '3C50FF', 'FFFFFF'),
    'A': ('arrival (the lab gate drops you here)', 'FF3C3C', 'FFFFFF'),
    'V': ('end portal (to the labkeeper)', '78DCFF', '000000'),
    'BAR': ('barkeeper', '28C83C', '000000'),
    'MW': ('weapon merchant', '28C83C', '000000'),
    'MA': ('armor merchant', '28C83C', '000000'),
}
SECTION_DEFAULTS = [('Atrium', 23, 24, 57, 69), ('Library', 59, 30, 93, 69), ('Crypt', 95, 1, 118, 98), ('Cistern', 23, 70, 91, 98)]
KINDS = {  # section number -> monster kinds (base template, name, scale %), used in turn
    1: [(423, 'Drowned Sailor', 112)],
    2: [(997, 'Gargoyle Archivist', 103)],
    3: [(423, 'Crypt Lord', 115), (1154, 'Drowned Knight', 106)],
    4: [(1158, 'Sea Golem', 106), (1150, 'Marsh Shaman', 109)],
}
WATERY = ('|', 's', 'S1', 'S2', 'S3', 'S4', 'E1', 'E2', 'E3', 'Q')


# ---------------------------------------------------------------------------- export
def export(path):
    import openpyxl
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.utils import get_column_letter
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    import draft_layout as d
    conv = {'k': 'Q', 'K': 'BAR', 'W': 'MW', 'R': 'MA', 'P': 'V', 'c': 'ct', '1': 'E1', '2': 'E2', '3': 'E3'}
    grid = [[None]*W for _ in range(H)]
    for y in range(H):
        for x in range(W):
            c = d.g[y][x]
            if c in 'SL':
                c = c+str([i+1 for i, s in enumerate(d.SECTIONS) if (s[4] if c == 'S' else s[1]) == (x, y)][0])
            elif c == 'D':
                c = 'D'+str([i+1 for i, s in enumerate(d.SECTIONS) if s[2] == (x, y)][0])
            else:
                c = conv.get(c, c)
            grid[y][x] = c

    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = 'Map'
    f = Font(name='Arial', size=8); fb = Font(name='Arial', size=8, bold=True)
    thin = Side(style='thin', color='666666')
    ws.cell(1, 1, 'y \\ x').font = fb
    for x in range(W):
        c = ws.cell(1, x+2, OX+x); c.font = fb; c.alignment = Alignment(horizontal='center', text_rotation=90)
        ws.column_dimensions[get_column_letter(x+2)].width = 3.6
    ws.column_dimensions['A'].width = 6
    ws.row_dimensions[1].height = 26
    for y in range(H):
        c = ws.cell(y+2, 1, OY+y); c.font = fb
        ws.row_dimensions[y+2].height = 18
        for x in range(W):
            c = ws.cell(y+2, x+2, grid[y][x]); c.font = f
            c.alignment = Alignment(horizontal='center', vertical='center')
            c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.freeze_panes = 'B2'
    rng = 'B2:%s%d' % (get_column_letter(W+1), H+1)
    for code, (desc, bg, fg) in CODES.items():
        ws.conditional_formatting.add(rng, FormulaRule(
            formula=['UPPER(B2)="%s"' % code.upper()],
            fill=PatternFill('solid', start_color=bg, end_color=bg, bgColor=bg),
            font=Font(name='Arial', size=8, color=fg), stopIfTrue=True))

    ws2 = wb.create_sheet('Legend')
    ws2['A1'] = 'Lab 11, the Sunken Dark - layout sheet'; ws2['A1'].font = Font(name='Arial', size=12, bold=True)
    notes = [
        'How to edit: type a code from the table below into any cell of the Map sheet. Codes are not case-sensitive; the colours follow automatically.',
        'The Map sheet covers map x %d-%d and y %d-%d (the headers show the real coordinates). Keep "#" everywhere outside the lab.' % (OX, OX+W-1, OY, OY+H-1),
        'Each of the 4 sections needs exactly one tome shelf (S1-S4), one lectern (L1-L4) and one gate (D1-D4) with its number.',
        'All water must lie inside one of the rectangles on the Sections sheet - that is what drains when the section\'s lectern gets its tome.',
        'Monsters (m) get their kind from the section they stand in. Put them on dry islands; water is no-monster ground.',
        'Exactly one each of A, X, V, Q, H, BAR, MW, MA, E1, E2, E3. If you move A, tell Claude (the server code has the arrival spot).',
        'When you are done, save the file and tell Claude: it checks the layout (swims, reachability) and rebuilds the lab from it.',
    ]
    for i, t in enumerate(notes):
        ws2.cell(3+i, 1, t).font = Font(name='Arial', size=10)
    ws2.cell(11, 1, 'Code').font = Font(name='Arial', size=10, bold=True)
    ws2.cell(11, 2, 'Meaning').font = Font(name='Arial', size=10, bold=True)
    for i, (code, (desc, bg, fg)) in enumerate(CODES.items()):
        c = ws2.cell(12+i, 1, code)
        c.fill = PatternFill('solid', start_color=bg, end_color=bg)
        c.font = Font(name='Arial', size=10, bold=True, color=fg); c.alignment = Alignment(horizontal='center')
        ws2.cell(12+i, 2, desc).font = Font(name='Arial', size=10)
    ws2.column_dimensions['A'].width = 8; ws2.column_dimensions['B'].width = 70

    ws3 = wb.create_sheet('Sections')
    for i, h in enumerate(('Section', 'Name', 'Water from x', 'from y', 'to x', 'to y')):
        ws3.cell(1, i+1, h).font = Font(name='Arial', size=10, bold=True)
    for i, (name, x0, y0, x1, y1) in enumerate(SECTION_DEFAULTS):
        for j, v in enumerate((i+1, name, OX+x0, OY+y0, OX+x1, OY+y1)):
            ws3.cell(2+i, j+1, v).font = Font(name='Arial', size=10, color='0000FF' if j >= 2 else '000000')
    ws3.cell(7, 1, 'Blue cells can be edited: the rectangle (map coordinates) whose water drains with that section\'s lectern. Rectangles must not overlap.').font = Font(name='Arial', size=10)
    for col in 'ABCDEF':
        ws3.column_dimensions[col].width = 14
    wb.save(path)
    print('wrote', path)


# ---------------------------------------------------------------------------- read + check
def read(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb['Map']
    up = {k.upper(): k for k in CODES}
    g = [['#']*W for _ in range(H)]
    errors = []
    for y in range(H):
        for x in range(W):
            v = ws.cell(y+2, x+2).value
            v = '#' if v is None or str(v).strip() == '' else str(v).strip()
            if v.upper() not in up:
                errors.append('unknown code "%s" at %d,%d' % (v, OX+x, OY+y)); v = '#'
            g[y][x] = up[v.upper()]
    secs = []
    for r in range(4):
        row = [wb['Sections'].cell(2+r, c+1).value for c in range(6)]
        secs.append((str(row[1]), (int(row[2])-OX, int(row[3])-OY, int(row[4])-OX, int(row[5])-OY)))
    return g, secs, errors


def build(path, outdir):
    from PIL import Image, ImageDraw
    random.seed(11)
    g, secs, errors = read(path)
    find = lambda code: [(x, y) for y in range(H) for x in range(W) if g[y][x] == code]

    def one(code):
        p = find(code)
        if len(p) != 1:
            errors.append('%s: need exactly one, found %d' % (code, len(p)))
            return p[0] if p else (0, 0)
        return p[0]
    SECTIONS = []   # name, lectern, gate, water rectangle, tome shelf
    for i, (name, r) in enumerate(secs):
        SECTIONS.append((name, one('L%d' % (i+1)), one('D%d' % (i+1)), r, one('S%d' % (i+1))))
    for code in ('A', 'X', 'V', 'Q', 'H', 'BAR', 'MW', 'MA', 'E1', 'E2', 'E3'):
        one(code)
    inrect = lambda x, y, r: r[0] <= x <= r[2] and r[1] <= y <= r[3]

    def water(x, y):    # is this tile water (for objects: what they stand on)
        c = g[y][x]
        if c == '~': return True
        if c not in WATERY: return False
        n = sum(1 for dx, dy in ((1, 0), (-1, 0), (0, 1), (0, -1)) if 0 <= x+dx < W and 0 <= y+dy < H and g[y+dy][x+dx] == '~')
        return n >= 2
    for y in range(H):
        for x in range(W):
            if water(x, y) and sum(inrect(x, y, s[3]) for s in SECTIONS) != 1:
                errors.append('water at %d,%d is not inside exactly one section rectangle' % (OX+x, OY+y))
    if errors:
        print('PROBLEMS:')
        for e in errors[:40]: print('  '+e)
        if len(errors) > 40: print('  ... and %d more' % (len(errors)-40))
        sys.exit(1)
    os.makedirs(outdir, exist_ok=True)
    json.dump([' '.join(r) for r in g], open(os.path.join(outdir, 'layout.json'), 'w'))
    ax, ay = one('A')
    print('arrival at %d,%d (server/data.h LAB11_X/LAB11_Y must say the same)' % (OX+ax, OY+ay))

    # ---- reachability and the longest swim on the best path
    walk = {'.', '~', 'e', 'A', 'T', 'E1', 'E2', 'E3', 'Q', 'D1', 'D2', 'D3', 'D4', 'SF', 'SD'}

    def swim(start, goals, gates_open):
        best = {}; dq = deque([(start, 0, 0)]); worst = None
        while dq:
            (x, y), run, mx = dq.popleft()
            if best.get((x, y), 1 << 30) <= mx: continue
            best[(x, y)] = mx
            if (x, y) in goals: worst = mx if worst is None else min(worst, mx)
            for dx, dy in ((1, 0), (-1, 0), (0, 1), (0, -1)):
                nx, ny = x+dx, y+dy
                if not (0 <= nx < W and 0 <= ny < H): continue
                c = g[ny][nx]
                if c[0] == 'D' and len(c) == 2 and (nx, ny) not in gates_open: continue
                if c not in walk and (nx, ny) not in goals: continue
                r = run+1 if water(nx, ny) else 0
                dq.append(((nx, ny), r, max(mx, r)))
        return worst
    adj = lambda p: {(p[0]+dx, p[1]+dy) for dx, dy in ((1, 0), (-1, 0), (0, 1), (0, -1))}
    gates = set(); start = (ax, ay); ok = True
    for name, lec, gate, r, book in SECTIONS:
        a = swim(start, adj(book), gates); b = swim(start, adj(lec), gates)
        print('%-8s tome: longest swim %s, lectern: %s' % (name, a, b))
        ok = ok and a is not None and b is not None
        gates.add(gate)
    v = swim(start, adj(one('V')), gates)
    print('end portal reachable:', v is not None); ok = ok and v is not None
    for c in ('E1', 'E2', 'E3', 'Q', 'BAR', 'MW', 'MA', 'X', 'H'):
        p = one(c); s = swim(start, adj(p) | {p}, gates)
        print('%-3s at %d,%d: %s' % (c, OX+p[0], OY+p[1], 'NOT REACHABLE' if s is None else 'swim %d' % s))
        ok = ok and s is not None
    if not ok: print('WARNING: something cannot be reached')

    # ---- preview
    S = 8; im = Image.new('RGB', (W*S, H*S)); d = ImageDraw.Draw(im)
    for y in range(H):
        for x in range(W):
            bg = CODES[g[y][x]][1]
            d.rectangle((x*S, y*S, x*S+S-1, y*S+S-1), fill=tuple(int(bg[i:i+2], 16) for i in (0, 2, 4)))
    im.save(os.path.join(outdir, 'preview.png'))

    # ---- build spec
    T_SHRINE, T_POLE, T_BOOK, T_SHELFB, T_SHELFE, T_LECT, T_GATE, T_BRAZ, T_PORTAL = 1330, 1331, 1334, 1338, 1342, 1343, 1347, 1348, 1349
    T_SHELFE2 = 1350     # 1350-1353: empty shelves in the four ordinary shelf looks
    C_KEEPER, C_BAR, C_WEAP, C_ARM, C_MON = 1200, 1201, 1202, 1203, 1210
    out = []; o = out.append
    o('# Lab 11, the Sunken Dark - generated by tools/world/lab11/lab11.py from %s (map %d,%d - %d,%d)' % (os.path.basename(path), OX, OY, OX+W-1, OY+H-1))
    o('clear %d %d %d %d' % (OX, OY, OX+W-1, OY+H-1))     # rebuilds start from an empty area
    o('# ---- item templates')
    o('it %d copy 673' % T_SHRINE); o('it %d name Shrine in Sunken Dark' % T_SHRINE); o('it %d d 0 2097152' % T_SHRINE)
    for i, bit in enumerate((262144, 524288, 1048576)):
        t = T_POLE+i
        o('it %d copy 1316' % t); o('it %d name Explotation Point 0/0/%d' % (t, bit)); o('it %d d 2 %d' % (t, bit)); o('it %d d 4 50000' % t)
    vols = ('I', 'II', 'III', 'IV'); bookspr = (1061, 1067, 1073, 1078)
    for i in range(4):
        t = T_BOOK+i
        o('it %d copy %d' % (t, bookspr[i])); o('it %d name Drowned Tome, Volume %s' % (t, vols[i])); o('it %d ref a drowned tome' % t)
        o('it %d desc A swollen, dripping book. Its pages still glow faintly. It belongs on a lectern.' % t)
        s = T_SHELFB+i
        o('it %d copy 1060' % s); o('it %d name Bookshelf' % s); o('it %d desc A waterlogged bookshelf. One book still looks whole.' % s)
        o('it %d d 0 %d' % (s, t)); o('it %d dur %d' % (s, TICKS*60))
    o('it %d copy 1060' % T_SHELFE); o('it %d name Bookshelf' % T_SHELFE); o('it %d desc A waterlogged bookshelf, its books swollen into pulp.' % T_SHELFE)
    o('it %d flags 613' % T_SHELFE); o('it %d driver 72' % T_SHELFE); o('it %d d 0 2' % T_SHELFE); o('it %d dur 0' % T_SHELFE)
    for j, spr in enumerate((16897, 16898, 16908, 16909)):   # the empty ones look like ordinary full shelves
        o('it %d copy %d' % (T_SHELFE2+j, T_SHELFE)); o('it %d spr %d %d' % (T_SHELFE2+j, spr, spr))
    for i, (name, lec, gate, r, book) in enumerate(SECTIONS):
        t = T_LECT+i
        o('it %d copy 1052' % t); o('it %d name Lectern' % t); o('it %d ref the lectern' % t)
        o('it %d desc A lectern of swollen wood, its top shaped for a large book.' % t)
        o('it %d spr 16895 16918' % t); o('it %d flags 1000000611' % t); o('it %d driver 72' % t); o('it %d light 0 0' % t); o('it %d dur 0' % t)
        vals = [1, T_BOOK+i, (OX+gate[0])+(OY+gate[1])*1024, OX+r[0], OY+r[1], OX+r[2], OY+r[3], TICKS*600, 0, 0]
        for j, v in enumerate(vals): o('it %d d %d %d' % (t, j, v))
    o('it %d copy 732' % T_GATE); o('it %d name Sunken Gate' % T_GATE); o('it %d ref the gate' % T_GATE)
    o('it %d desc Rusted bars. There is no lock - the water behind it holds it shut.' % T_GATE)
    o('it %d flags 611' % T_GATE); o('it %d driver 72' % T_GATE); o('it %d dur 0' % T_GATE)
    for j in range(10): o('it %d d %d %d' % (T_GATE, j, 3 if j == 0 else 0))
    o('it %d copy 741' % T_BRAZ); o('it %d name Brazier' % T_BRAZ); o('it %d ref the brazier' % T_BRAZ)
    o('it %d desc A cold brazier, still holding a little dry coal.' % T_BRAZ)
    o('it %d spr 18129 16510' % T_BRAZ); o('it %d status 0 16' % T_BRAZ); o('it %d flags 400211' % T_BRAZ)
    o('it %d light 0 175' % T_BRAZ); o('it %d dur %d' % (T_BRAZ, TICKS*300))
    o('it %d copy 1299' % T_PORTAL); o('it %d name Magic Portal (Lab 11 End)' % T_PORTAL)
    for j, v in enumerate((33, 597, 11, 210000)): o('it %d d %d %d' % (T_PORTAL, j, v))

    o('# ---- character templates')
    def chbase(t, frm, name, x, y, scale=None, monster=True):
        o('ch %d copy %d' % (t, frm)); o('ch %d name %s' % (t, name))
        for i in list(range(10, 24))+[31, 32, 33, 34, 35, 60, 73]: o('ch %d d %d 0' % (t, i))
        o('ch %d pos %d %d' % (t, x, y))
        if scale: o('ch %d scale %d' % (t, scale))
        if monster:
            for i, v in ((42, 1101), (43, 1101), (44, 1), (45, 0), (46, 0), (59, 1101)): o('ch %d d %d %d' % (t, i, v))
            o('ch %d setflags 10000' % t)      # sees in the dark
    chbase(C_KEEPER, 423, 'Drowned Lord', 0, 0, 115)
    o('ch %d ref the Drowned Lord' % C_KEEPER); o('ch %d clrflags 10' % C_KEEPER)
    x, y = one('BAR'); chbase(C_BAR, 1083, 'Barkeeper', OX+x, OY+y, monster=False)
    x, y = one('MW'); chbase(C_WEAP, 1091, 'Merchant', OX+x, OY+y, monster=False)
    o('ch %d items 28 32 36 29 33 37 30 34 38 523 524 125' % C_WEAP)
    x, y = one('MA'); chbase(C_ARM, 1092, 'Merchant', OX+x, OY+y, monster=False)
    o('ch %d items 63 64 65 66 67 68 69 70 71 72 73 74 981 982 983 984 985 986 75 76 77 78 79 80' % C_ARM)
    t = C_MON; count = {}
    for x, y in find('m'):
        sec = ([i+1 for i, s in enumerate(SECTIONS) if inrect(x, y, s[3])] or [4])[0]
        k = KINDS[sec][count.get(sec, 0) % len(KINDS[sec])]; count[sec] = count.get(sec, 0)+1
        chbase(t, k[0], k[1], OX+x, OY+y, k[2]); t += 1
    print('monsters per section:', dict(sorted(count.items())), 'templates %d-%d' % (C_MON, t-1))
    if t > 1260: print('WARNING: more than 50 monsters - templates run past 1259')

    o('# ---- tiles')
    SEAWEED = (340, 342, 348, 350, 352, 354)
    FL_IN, FL_UW, FL_NOM, FL_TAV, FL_MB, FL_SB = 0x4, 0x8, 0x20, 0x80, 0x1, 0x2
    hall = {'e', 'A', 'BAR', 'MW', 'MA', 'X', 'H', 'T', 'ct'}
    lect = {s[1]: i for i, s in enumerate(SECTIONS)}; shelf = {s[4]: i for i, s in enumerate(SECTIONS)}
    items = []
    for y in range(H):
        for x in range(W):
            c = g[y][x]; X, Y = OX+x, OY+y
            nb = [g[y+dy][x+dx] for dx in (-1, 0, 1) for dy in (-1, 0, 1) if 0 <= x+dx < W and 0 <= y+dy < H]
            if c == '#':
                if any(n != '#' for n in nb): o('tile %d %d 1100 1214 %x' % (X, Y, FL_MB | FL_SB | FL_IN | FL_NOM))
                else: o('tile %d %d 1008 0 %x' % (X, Y, FL_MB | FL_SB | FL_IN | FL_NOM))   # indoors too, or daylight lights the rock
            elif water(x, y):
                fs = random.choice(SEAWEED) if c == '~' and random.random() < 0.04 else 0
                o('tile %d %d 1008 %d %x' % (X, Y, fs, FL_IN | FL_UW | FL_NOM))
            elif c == 'SW':
                o('tile %d %d 1100 556 %x' % (X, Y, FL_MB | FL_SB | FL_IN | FL_NOM))
            elif c in ('SF', 'SC', 'SD'):
                o('tile %d %d 558 0 %x' % (X, Y, FL_IN | FL_NOM))
            elif c == 'ct':
                o('tile %d %d 558 176 %x' % (X, Y, FL_IN | FL_NOM | FL_MB))
            elif c in hall or (c == 'C' and any(n in hall for n in nb)):
                o('tile %d %d 558 0 %x' % (X, Y, FL_IN | FL_NOM | (FL_TAV if c == 'T' else 0)))
            else:   # dry stone floor; no monsters on landings by gates, lecterns and the vault portal
                nom = c[0] in 'DLV' or any(n in hall for n in nb)
                o('tile %d %d 1100 0 %x' % (X, Y, FL_IN | (FL_NOM if nom else 0)))
            tmp = {'B': T_BRAZ, 'C': 741, 'V': T_PORTAL, 'X': 661, 'H': 603, 'Q': T_SHRINE,
                   'E1': T_POLE, 'E2': T_POLE+1, 'E3': T_POLE+2, 'SD': 559, 'SC': 296}.get(c)
            if c == 's': tmp = T_SHELFE2+random.randrange(4)
            if c in ('D1', 'D2', 'D3', 'D4'): tmp = T_GATE
            if c == '|': tmp = random.choice((1053, 1054, 1058, 1059))
            if c in ('L1', 'L2', 'L3', 'L4'): tmp = T_LECT+lect[(x, y)]
            if c in ('S1', 'S2', 'S3', 'S4'): tmp = T_SHELFB+shelf[(x, y)]
            if tmp: items.append('place %d %d %d' % (tmp, X, Y))
    out += items
    out += ['movenpc %d' % t for t in (C_BAR, C_WEAP, C_ARM)]   # the inn NPCs walk to their (new) places
    open(os.path.join(outdir, 'lab11.spec'), 'w').write('\n'.join(out)+'\n')
    print('spec lines:', len(out), 'items:', len(items))


if __name__ == '__main__':
    if len(sys.argv) >= 3 and sys.argv[1] == 'export': export(sys.argv[2])
    elif len(sys.argv) >= 4 and sys.argv[1] == 'build': build(sys.argv[2], sys.argv[3])
    else: print(__doc__)
