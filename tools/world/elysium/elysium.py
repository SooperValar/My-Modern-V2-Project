#!/usr/bin/env python3
"""Elysium (the gods' area) - layout sheet.

    elysium.py export <dump.tsv> <xlsx>    write the sheet from the live map (area_dump output)
    elysium.py diff   <xlsx>               list what was changed against the export

dump.tsv: area_dump <.dat folder> 771 829 788 826 > dump.tsv
          (x y sprite fsprite flags itemtemp itemname itemsprite driver chartemp charname charsprite)

Sheet "Map": one cell per tile, orientation like the minimap: row = map x (north at the top),
column = map y (west on the left). The hidden sheet "Original" keeps the exported codes, so a
build only touches the cells that were changed; everything left alone keeps its exact sprites.
"""
import sys
X0, X1, Y0, Y1 = 771, 829, 788, 826
MERCH = {1180: 'Armsmaster', 1181: 'Beltmaker', 1182: 'Armorer', 1183: 'Amulet Keeper', 1184: 'Ringmaker',
         1185: 'Keymaster', 1186: 'Alchemist', 1187: 'Questkeeper', 1188: 'Curator', 1189: 'Scribe'}
CODES = {  # code: (meaning, fill, text)
    '': ('outside (nothing - leave as it is)', '101014', 'FFFFFF'),
    '#': ('wall', '2B2B30', 'FFFFFF'),
    '.': ('floor', 'C8C2B4', '000000'),
    'F': ('fire bowl (light)', 'FFA028', '000000'),
    'D': ('door', 'C87828', '000000'),
    'H': ('shrine', 'F0DC3C', '000000'),
    'T': ('trash can (destroys what is put in)', '7F7F7F', 'FFFFFF'),
    'S': ('shield on the wall (decoration)', 'A0A0D0', '000000'),
    'i': ('other item (stays where it is)', 'D2B48C', '000000'),
    'n': ('display NPC (the monster gallery)', 'C81E1E', 'FFFFFF'),
    't': ('table (blocks walking, not sight)', '8B5A2B', 'FFFFFF'),
    'c': ('column / pillar', 'E8E0C8', '000000'),
    'g': ('gate or arch piece (dark iron)', '4A4A60', 'FFFFFF'),
    'P': ('portal (a new one needs a ?note saying where it goes)', '3C50FF', 'FFFFFF'),
}
for k, (t, nm) in enumerate(MERCH.items()):
    CODES['M%d' % (k+1)] = ('god shop merchant: %s (template %d)' % (nm, t), '28C83C', '000000')
NAMES = {'Fire in Bowl': 'F', 'Black Stronghold Iron Door': 'D', 'Shrine': 'H', 'Trash Can': 'T', 'Shield': 'S'}


def load(tsv):
    T = {}
    for l in open(tsv):
        f = l.rstrip('\n').split('\t')
        T[int(f[0]), int(f[1])] = f
    return T


def code(f):
    spr, fs, fl = int(f[2]), int(f[3]), int(f[4], 16)
    if f[9]:
        t = int(f[9])
        if t in MERCH:
            return 'M%d' % (t-1179)
        return 'n'
    if f[5]:
        if 'Portal' in f[6] and 'floor tile' not in f[6]:
            return 'P'
        return NAMES.get(f[6], 'i')
    if spr == 1008 and not fs:
        return ''
    if fs in (1242, 1244):
        return 't'
    if fs == 106:
        return 'c'
    if fs in (968, 970, 976, 978):
        return 'g'
    if fs or fl & 1:
        return '#'
    return '.'


def export(tsv, path):
    import openpyxl
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.utils import get_column_letter
    T = load(tsv)
    W, H = Y1-Y0+1, X1-X0+1
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = 'Map'
    org = wb.create_sheet('Original')
    f = Font(name='Arial', size=8); fb = Font(name='Arial', size=8, bold=True)
    thin = Side(style='thin', color='666666')
    ws.cell(1, 1, 'x \\ y').font = fb
    for j in range(W):
        c = ws.cell(1, j+2, Y0+j); c.font = fb; c.alignment = Alignment(horizontal='center', text_rotation=90)
        ws.column_dimensions[get_column_letter(j+2)].width = 3.6
    ws.column_dimensions['A'].width = 6
    ws.row_dimensions[1].height = 26
    objs = []
    for i in range(H):
        x = X0+i
        c = ws.cell(i+2, 1, x); c.font = fb
        ws.row_dimensions[i+2].height = 18
        for j in range(W):
            y = Y0+j
            fl = T[x, y]
            v = code(fl)
            c = ws.cell(i+2, j+2, v or None); c.font = f
            c.alignment = Alignment(horizontal='center', vertical='center')
            c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            org.cell(i+2, j+2, v or None)
            if fl[9]:
                objs.append((x, y, 'NPC', fl[10], int(fl[9]), v))
            elif fl[5]:
                objs.append((x, y, 'item', fl[6], int(fl[5]), v))
    org.sheet_state = 'hidden'
    ws.freeze_panes = 'B2'
    rng = 'B2:%s%d' % (get_column_letter(W+1), H+1)
    ws.conditional_formatting.add(rng, FormulaRule(formula=['LEFT(B2,1)="?"'],
        fill=PatternFill('solid', start_color='FF3CFF', end_color='FF3CFF', bgColor='FF3CFF'),
        font=Font(name='Arial', size=8, bold=True, color='000000'), stopIfTrue=True))
    ws.conditional_formatting.add(rng, FormulaRule(formula=['ISBLANK(B2)'],
        fill=PatternFill('solid', start_color='101014', end_color='101014', bgColor='101014'), stopIfTrue=True))
    for k, (desc, bg, fg) in CODES.items():
        if not k:
            continue
        ws.conditional_formatting.add(rng, FormulaRule(
            formula=['EXACT(B2,"%s")' % k] if k in ('i', 'n', 'S', 's', 't', 'T', 'c', 'g') else ['UPPER(B2)="%s"' % k.upper()],
            fill=PatternFill('solid', start_color=bg, end_color=bg, bgColor=bg),
            font=Font(name='Arial', size=8, color=fg), stopIfTrue=True))

    ws2 = wb.create_sheet('Legend', 1)
    ws2['A1'] = 'Elysium - layout sheet'; ws2['A1'].font = Font(name='Arial', size=12, bold=True)
    notes = [
        'The Map sheet is the live Elysium as built, map x %d-%d (rows, north at the top) and y %d-%d (columns, west on the left) - the same way round as the minimap.' % (X0, X1, Y0, Y1),
        'How to edit: type a code from the table below into a cell. The colours follow automatically. Only cells you change are rebuilt; the rest keeps its exact look.',
        'New wall or floor copies the style of the walls / floor next to it. Make new rooms by painting "." with "#" around them over the dark outside cells.',
        'Move a merchant: type its code (M1-M10) where it should stand and clear (or set to ".") its old cell.',
        'Anything else you want (a portal, a sign, a bank, a chest, an NPC...): type ? and a short note into the cell, e.g. "?portal to Aston" or "?sign: Gods only". It shows pink, and Claude builds it.',
        'The zone name "in the Elysium" currently covers only x 792-813, y 796-811 (the middle hall); it will be widened to everything you keep.',
        'When you are done, save the file and tell Claude.',
    ]
    for i, t in enumerate(notes):
        ws2.cell(3+i, 1, t).font = Font(name='Arial', size=10)
    ws2.cell(11, 1, 'Code').font = Font(name='Arial', size=10, bold=True)
    ws2.cell(11, 2, 'Meaning').font = Font(name='Arial', size=10, bold=True)
    r = 12
    for k, (desc, bg, fg) in list(CODES.items()) + [('?...', ('your note - Claude builds what it says', 'FF3CFF', '000000'))]:
        c = ws2.cell(r, 1, k if k else '(empty)')
        c.fill = PatternFill('solid', start_color=bg, end_color=bg)
        c.font = Font(name='Arial', size=10, bold=True, color=fg); c.alignment = Alignment(horizontal='center')
        ws2.cell(r, 2, desc).font = Font(name='Arial', size=10)
        r += 1
    ws2.column_dimensions['A'].width = 9; ws2.column_dimensions['B'].width = 110

    ws3 = wb.create_sheet('Objects', 2)
    for i, h in enumerate(('x', 'y', 'kind', 'name', 'template', 'code')):
        ws3.cell(1, i+1, h).font = Font(name='Arial', size=10, bold=True)
    for r, o in enumerate(sorted(objs, key=lambda o: (o[2], o[0], o[1]))):
        for i, v in enumerate(o):
            ws3.cell(r+2, i+1, v).font = Font(name='Arial', size=10)
    ws3.column_dimensions['D'].width = 30
    wb.save(path)
    print('wrote', path, len(objs), 'objects')


def read(path, sheet):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet]
    g = {}
    for i in range(X1-X0+1):
        for j in range(Y1-Y0+1):
            v = ws.cell(i+2, j+2).value
            g[X0+i, Y0+j] = '' if v is None else str(v).strip()
    return g


def diff(path):
    a, b = read(path, 'Original'), read(path, 'Map')
    for k in sorted(a):
        if a[k] != b[k]:
            print('%d,%d: %r -> %r' % (k[0], k[1], a[k], b[k]))


if __name__ == '__main__':
    if sys.argv[1] == 'export':
        export(sys.argv[2], sys.argv[3])
    elif sys.argv[1] == 'diff':
        diff(sys.argv[2])
